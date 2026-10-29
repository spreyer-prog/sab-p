import base64
import io

from openpyxl import Workbook

from odoo.tests.common import TransactionCase


class TestSabCalculationImportReferences(TransactionCase):
    """Import 0.35 as the payable 35 %, i.e. a 65 % discount."""

    @staticmethod
    def _workbook_bytes(discount=0.35):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Kalkulation"
        row = 21
        worksheet.cell(row, 4, 0.98)   # Mechanikfaktor
        worksheet.cell(row, 5, 0.00)   # Verdrahtungsfaktor
        worksheet.cell(row, 6, 0.02)   # Prueffaktor
        worksheet.cell(row, 8, 83.50)  # Material-Listenpreis
        worksheet.cell(row, 9, 10)     # VE
        if discount is not None:
            worksheet.cell(row, 10, discount)  # Preisfaktor, kein Rabatt-Prozent
        worksheet.cell(row, 11, 0)     # Zeit
        worksheet.cell(row, 12, 0)     # Hilfsmaterial
        worksheet.cell(row, 18, "Plantasche DIN A4 – Import-Test")
        worksheet.cell(row, 19, 4.19)  # Excel VK als Referenz
        worksheet.cell(row, 43, 0)
        output = io.BytesIO()
        workbook.save(output)
        return base64.b64encode(output.getvalue())

    def _run_import(self, discount=0.35):
        wizard = self.env["sab.calculation.import"].create({
            "filename": "referenz.xlsx",
            "file_data": self._workbook_bytes(discount),
        })
        wizard.action_check_file()
        wizard.action_import()
        return wizard

    def test_discount_factor_035_means_65_percent_discount(self):
        wizard = self._run_import(0.35)
        item = self.env["sab.calculation.item"].search([
            ("name", "=", "Plantasche DIN A4 – Import-Test")
        ], limit=1)
        self.assertTrue(item)
        self.assertEqual(wizard.imported_count, 1)
        self.assertAlmostEqual(item.legacy_list_price, 83.5)
        self.assertAlmostEqual(item.legacy_packaging_unit, 10.0)
        self.assertAlmostEqual(item.legacy_discount_factor, 0.35)
        self.assertAlmostEqual(item.legacy_discount_percent, 65.0)
        self.assertAlmostEqual(item.legacy_net_purchase_unit, 2.9225)
        self.assertAlmostEqual(item.legacy_sales_price, 4.19)
        # Excel prices are references, not live supplier prices.
        self.assertFalse(item.product_line_ids)
        self.assertAlmostEqual(item.purchase_total, 0.0)

    def test_reimport_enriches_existing_item_without_creating_duplicates(self):
        self._run_import(0.35)
        item = self.env["sab.calculation.item"].search([
            ("name", "=", "Plantasche DIN A4 – Import-Test")
        ], limit=1)
        item.legacy_discount_factor = 1.0
        second = self._run_import(0.35)
        self.assertEqual(second.imported_count, 0)
        self.assertEqual(second.updated_count, 1)
        self.assertAlmostEqual(item.legacy_discount_factor, 0.35)
        self.assertEqual(self.env["sab.calculation.item"].search_count([
            ("name", "=", "Plantasche DIN A4 – Import-Test")
        ]), 1)

    def test_blank_discount_is_no_discount(self):
        self._run_import(None)
        item = self.env["sab.calculation.item"].search([
            ("name", "=", "Plantasche DIN A4 – Import-Test")
        ], limit=1)
        self.assertAlmostEqual(item.legacy_discount_factor, 1.0)
        self.assertAlmostEqual(item.legacy_discount_percent, 0.0)
        self.assertAlmostEqual(item.legacy_net_purchase_unit, 8.35)

    def test_explicit_zero_multiplier_is_preserved(self):
        self._run_import(0.0)
        item = self.env["sab.calculation.item"].search([
            ("name", "=", "Plantasche DIN A4 – Import-Test")
        ], limit=1)
        self.assertAlmostEqual(item.legacy_discount_factor, 0.0)
        self.assertAlmostEqual(item.legacy_net_purchase_unit, 0.0)
