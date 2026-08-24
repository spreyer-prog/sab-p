import base64
import hashlib
import io

from odoo import fields, models
from odoo.exceptions import UserError


class SabCalculationImport(models.TransientModel):
    _name = "sab.calculation.import"
    _description = "SAB-P Kalkulationsdatenbank importieren"

    file_data = fields.Binary(
        string="Excel-Datei",
        required=True,
        attachment=False,
    )

    filename = fields.Char(
        string="Dateiname",
    )

    state = fields.Selection(
        selection=[
            ("upload", "Datei auswählen"),
            ("checked", "Datei geprüft"),
            ("imported", "Import abgeschlossen"),
        ],
        string="Status",
        default="upload",
        readonly=True,
    )

    workbook_info = fields.Text(
        string="Prüfergebnis",
        readonly=True,
    )

    detected_sheet_count = fields.Integer(
        string="Anzahl Tabellenblätter",
        readonly=True,
    )

    detected_row_count = fields.Integer(
        string="Erkannte Zeilen",
        readonly=True,
    )

    imported_count = fields.Integer(
        string="Importierte Kalkulationsartikel",
        readonly=True,
    )

    skipped_count = fields.Integer(
        string="Übersprungene Zeilen",
        readonly=True,
    )

    duplicate_count = fields.Integer(
        string="Dubletten",
        readonly=True,
    )

    # ---------------------------------------------------------
    # Datei
    # ---------------------------------------------------------

    def _get_file_extension(self):
        self.ensure_one()

        filename = (self.filename or "").strip().lower()

        if filename.endswith(".xls"):
            return "xls"

        if filename.endswith(".xlsx"):
            return "xlsx"

        return False

    def _get_file_content(self):
        self.ensure_one()

        if not self.file_data:
            raise UserError(
                "Bitte zuerst die SAB-P Kalkulationsdatenbank auswählen."
            )

        try:
            content = base64.b64decode(self.file_data)
        except Exception as exc:
            raise UserError(
                "Die hochgeladene Datei konnte nicht verarbeitet werden."
            ) from exc

        if not content:
            raise UserError("Die hochgeladene Datei ist leer.")

        return content

    # ---------------------------------------------------------
    # Hilfsfunktionen
    # ---------------------------------------------------------

    @staticmethod
    def _clean_text(value):
        if value is None:
            return ""

        text = str(value)
        text = text.replace("\r\n", "\n")
        text = text.replace("\r", "\n")
        text = text.strip()

        return text

    @staticmethod
    def _number_or_none(value):
        if value is None or value == "":
            return None

        if isinstance(value, bool):
            return None

        if isinstance(value, (int, float)):
            return float(value)

        try:
            text = str(value).strip()
            text = text.replace(".", "")
            text = text.replace(",", ".")
            return float(text)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _percent_value(value):
        """
        Die alte Excel-Datei speichert:
        0,98 = 98 %
        0,08 = 8 %
        0,02 = 2 %

        Odoo zeigt den Faktor als Prozentwert an:
        98 / 8 / 2
        """
        number = SabCalculationImport._number_or_none(value)

        if number is None:
            return 0.0

        return number * 100.0

    @staticmethod
    def _space_factor_value(value):
        """
        AQ wird NICHT prozentual umgerechnet.

        0,50 bleibt 0,50
        1,00 bleibt 1,00
        1,25 bleibt 1,25
        """
        number = SabCalculationImport._number_or_none(value)

        if number is None:
            return 0.0

        return number

    @staticmethod
    def _make_import_key(
        description,
        mechanical_factor,
        wiring_factor,
        testing_factor,
        space_factor,
    ):
        raw_key = "|".join([
            description.strip().lower(),
            f"{mechanical_factor:.6f}",
            f"{wiring_factor:.6f}",
            f"{testing_factor:.6f}",
            f"{space_factor:.6f}",
        ])

        return hashlib.sha1(
            raw_key.encode("utf-8")
        ).hexdigest()

    # ---------------------------------------------------------
    # XLS
    # ---------------------------------------------------------

    def _read_xls_sheets(self, file_content):
        try:
            import xlrd
        except ImportError as exc:
            raise UserError(
                "Die Python-Bibliothek 'xlrd' ist nicht installiert."
            ) from exc

        try:
            workbook = xlrd.open_workbook(
                file_contents=file_content,
                on_demand=True,
            )
        except Exception as exc:
            raise UserError(
                "Die XLS-Datei konnte nicht gelesen werden.\n\n"
                f"{exc}"
            ) from exc

        result = []

        for sheet_name in workbook.sheet_names():
            sheet = workbook.sheet_by_name(sheet_name)

            rows = []

            # Excel-Zeile 21 = Index 20
            for row_index in range(20, sheet.nrows):
                rows.append({
                    "excel_row": row_index + 1,
                    "D": sheet.cell_value(row_index, 3),
                    "E": sheet.cell_value(row_index, 4),
                    "F": sheet.cell_value(row_index, 5),
                    "R": sheet.cell_value(row_index, 17),
                    "AQ": (
                        sheet.cell_value(row_index, 42)
                        if sheet.ncols > 42
                        else None
                    ),
                })

            result.append({
                "name": sheet_name,
                "row_count": sheet.nrows,
                "column_count": sheet.ncols,
                "rows": rows,
            })

        workbook.release_resources()

        return result

    # ---------------------------------------------------------
    # XLSX
    # ---------------------------------------------------------

    def _read_xlsx_sheets(self, file_content):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise UserError(
                "Die Python-Bibliothek 'openpyxl' ist nicht installiert."
            ) from exc

        try:
            workbook = load_workbook(
                filename=io.BytesIO(file_content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise UserError(
                "Die XLSX-Datei konnte nicht gelesen werden.\n\n"
                f"{exc}"
            ) from exc

        result = []

        for worksheet in workbook.worksheets:
            rows = []

            for row_number in range(21, (worksheet.max_row or 0) + 1):
                rows.append({
                    "excel_row": row_number,
                    "D": worksheet.cell(row_number, 4).value,
                    "E": worksheet.cell(row_number, 5).value,
                    "F": worksheet.cell(row_number, 6).value,
                    "R": worksheet.cell(row_number, 18).value,
                    "AQ": worksheet.cell(row_number, 43).value,
                })

            result.append({
                "name": worksheet.title,
                "row_count": worksheet.max_row or 0,
                "column_count": worksheet.max_column or 0,
                "rows": rows,
            })

        workbook.close()

        return result

    def _read_workbook(self):
        self.ensure_one()

        extension = self._get_file_extension()

        if not extension:
            raise UserError(
                "Es werden nur .xls- und .xlsx-Dateien unterstützt."
            )

        file_content = self._get_file_content()

        if extension == "xls":
            return self._read_xls_sheets(file_content)

        return self._read_xlsx_sheets(file_content)

    # ---------------------------------------------------------
    # Datei prüfen
    # ---------------------------------------------------------

    def action_check_file(self):
        self.ensure_one()

        sheets = self._read_workbook()

        if not sheets:
            raise UserError(
                "In der Excel-Datei wurden keine Tabellenblätter gefunden."
            )

        total_rows = sum(
            sheet["row_count"]
            for sheet in sheets
        )

        info_lines = [
            "SAB-P Kalkulationsdatenbank erfolgreich gelesen.",
            "",
            f"Datei: {self.filename or '-'}",
            f"Tabellenblätter: {len(sheets)}",
            f"Zeilen gesamt: {total_rows}",
            "",
            "Verbindliche Importzuordnung:",
            "R  = Beschreibung / Angebotstext",
            "D  = Mechanikfaktor %",
            "E  = Verdrahtungsfaktor %",
            "F  = Prüffaktor %",
            "AQ = Platzfaktor (Dezimalwert)",
            "",
            "NICHT importiert werden:",
            "AJ = fertige Mechanikminuten",
            "AK = fertige Verdrahtungsminuten",
            "AL = fertige Prüfminuten",
            "AI = fertig berechnete Platzeinheiten",
            "",
            "20 Suchbegriff-Felder bleiben je Kalkulationsartikel frei verfügbar.",
            "",
            "Gefundene Tabellenblätter:",
        ]

        for sheet in sheets:
            info_lines.append(
                f"- {sheet['name']}: "
                f"{sheet['row_count']} Zeilen / "
                f"{sheet['column_count']} Spalten"
            )

        self.write({
            "state": "checked",
            "workbook_info": "\n".join(info_lines),
            "detected_sheet_count": len(sheets),
            "detected_row_count": total_rows,
            "imported_count": 0,
            "skipped_count": 0,
            "duplicate_count": 0,
        })

        return self._reopen_wizard()

    # ---------------------------------------------------------
    # Echter Import
    # ---------------------------------------------------------

    def action_import(self):
        self.ensure_one()

        if self.state != "checked":
            raise UserError(
                "Bitte die Datei vor dem Import zuerst prüfen."
            )

        sheets = self._read_workbook()

        calculation_model = self.env["sab.calculation.item"]

        imported_count = 0
        skipped_count = 0
        duplicate_count = 0

        keys_seen_in_file = set()

        for sheet in sheets:
            sheet_name = sheet["name"]

            for row in sheet["rows"]:
                description = self._clean_text(row["R"])

                # Ohne Beschreibung kein Kalkulationsartikel.
                if not description:
                    skipped_count += 1
                    continue

                mechanical_raw = self._number_or_none(row["D"])
                wiring_raw = self._number_or_none(row["E"])
                testing_raw = self._number_or_none(row["F"])
                space_raw = self._number_or_none(row["AQ"])

                # Reine Text-/Hinweiszeilen werden nicht als
                # Kalkulationsartikel importiert.
                if (
                    mechanical_raw is None
                    and wiring_raw is None
                    and testing_raw is None
                    and space_raw is None
                ):
                    skipped_count += 1
                    continue

                mechanical_factor = self._percent_value(row["D"])
                wiring_factor = self._percent_value(row["E"])
                testing_factor = self._percent_value(row["F"])
                space_factor = self._space_factor_value(row["AQ"])

                import_key = self._make_import_key(
                    description,
                    mechanical_factor,
                    wiring_factor,
                    testing_factor,
                    space_factor,
                )

                # Gleiche Position mehrfach auf verschiedenen
                # Excel-Blättern -> nur einmal übernehmen.
                if import_key in keys_seen_in_file:
                    duplicate_count += 1
                    continue

                keys_seen_in_file.add(import_key)

                # Schutz bei erneutem Import derselben Datei.
                existing = calculation_model.search(
                    [
                        ("legacy_import_key", "=", import_key),
                    ],
                    limit=1,
                )

                if existing:
                    duplicate_count += 1
                    continue

                calculation_model.create({
                    "name": description,
                    "quotation_text": description,
                    "mechanical_factor": mechanical_factor,
                    "wiring_factor": wiring_factor,
                    "testing_factor": testing_factor,
                    "space_factor": space_factor,
                    "legacy_import_key": import_key,
                    "legacy_source_sheet": sheet_name,
                    "legacy_source_row": row["excel_row"],
                })

                imported_count += 1

        result_text = [
            "Import abgeschlossen.",
            "",
            f"Neu importiert: {imported_count}",
            f"Dubletten übersprungen: {duplicate_count}",
            f"Sonstige Zeilen übersprungen: {skipped_count}",
            "",
            "Es wurden ausschließlich die vereinbarten "
            "Kalkulationshüllen übernommen.",
            "",
            "Keine fertigen Minutenwerte wurden importiert.",
            "Keine berechneten Platzeinheiten wurden importiert.",
            "Keine DATANORM-Artikel wurden zugeordnet.",
            "Die 20 Suchbegriffe bleiben zur späteren Pflege frei.",
        ]

        self.write({
            "state": "imported",
            "imported_count": imported_count,
            "skipped_count": skipped_count,
            "duplicate_count": duplicate_count,
            "workbook_info": "\n".join(result_text),
        })

        return self._reopen_wizard()

    # ---------------------------------------------------------
    # Zurücksetzen
    # ---------------------------------------------------------

    def action_reset(self):
        self.ensure_one()

        self.write({
            "state": "upload",
            "workbook_info": False,
            "detected_sheet_count": 0,
            "detected_row_count": 0,
            "imported_count": 0,
            "skipped_count": 0,
            "duplicate_count": 0,
        })

        return self._reopen_wizard()

    def _reopen_wizard(self):
        return {
            "type": "ir.actions.act_window",
            "name": "Kalkulationsdatenbank importieren",
            "res_model": "sab.calculation.import",
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
        }